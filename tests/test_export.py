"""Tests for witcher3.export. Run with: python -m unittest -v"""
import tempfile
import unittest
from pathlib import Path

import openpyxl

from tests.helpers import real
from witcher3.baselines import play
from witcher3.export import ORIGINAL, export
from witcher3.greedy import greedy
from witcher3.loader import restrict
from witcher3.simulator import Model
from witcher3.units import Plan

REAL = real()
HIDDEN = {'Before', 'After', 'Block_ID', 'Path_ID', 'During'}


class Export(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.scope = ('base', 'hos')
        cls.ds = restrict(REAL, cls.scope)
        cls.plan = Plan(Model(cls.ds))
        cls.result = play(cls.plan, greedy(cls.plan))
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = export(cls.ds, cls.plan, cls.result, Path(cls.tmp.name) / 'sub' / 'plan.xlsx', cls.scope, REAL.level_table)
        cls.wb = openpyxl.load_workbook(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_sheets(self):
        self.assertEqual(self.wb.sheetnames, ['Summary', 'Plan', 'Skipped'])

    def test_plan_has_every_original_column_and_none_of_the_search_columns(self):
        head = [c.value for c in self.wb['Plan'][1]]
        self.assertFalse(HIDDEN & set(head))
        for name, _ in ORIGINAL:
            self.assertIn(name, head)
        self.assertEqual(head[0], 'Order')

    def test_plan_rows_are_the_run_in_order(self):
        ws = self.wb['Plan']
        head = [c.value for c in ws[1]]
        ids = [r[head.index('ID')] for r in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(ids, self.result.order)
        total = [r[head.index('Total XP')] for r in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(total, sorted(total))                                  # XP never goes down
        self.assertEqual(ws.cell(ws.max_row, head.index('Level after') + 1).value, self.result.level)

    def test_original_values_are_kept(self):
        ws = self.wb['Plan']
        head = [c.value for c in ws[1]]
        row = [c.value for c in ws[2]]
        self.assertEqual(row[head.index('ID')], 'WOF1_S1')
        self.assertEqual(row[head.index('Quest_Name')], 'Bez i agrest')
        self.assertEqual(row[head.index('Step_ID')], 1)                         # numbers stay numbers
        self.assertEqual(row[head.index('Location')], 'Biały Sad')              # Polish text survives

    def test_skipped_plus_played_is_everything(self):
        n_skipped = self.wb['Skipped'].max_row - 1
        self.assertEqual(n_skipped + len(self.result.order), len(self.plan.model.uids))
        reasons = {r[-1] for r in self.wb['Skipped'].iter_rows(min_row=2, values_only=True)}
        self.assertNotIn('Never became available (check the data)', reasons)

    def test_no_formulas(self):
        for ws in self.wb:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    self.assertFalse(isinstance(v, str) and v.startswith('='), (ws.title, v))


if __name__ == '__main__':
    unittest.main()
