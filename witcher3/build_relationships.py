"""Regenerate the '<Location> Relationships' sheets from Before / After / Block_ID / Path_ID.

Usage: python -m witcher3.build_relationships [--file PATH]   (overwrites those sheets and Relationship_issues)

After  -> requires (one prerequisite) or unlocks (one row per alternative: 'S2a/b', 'A|B')
Before -> blocks
Block  -> in_block  (source = block tag, target = member step)
Path   -> excludes_path between the path tags of one decision (same quest, same outside prerequisites), and
          between tags listed together in one cell
Name-based cells (Polish quest names) resolve to the last mandatory step of the quest.
"""
import argparse
import collections
import itertools
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font

if __package__ in (None, ''):  # run as a plain script (python witcher3/build_relationships.py, an IDE Run button)
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from witcher3.loader import DATA_FILE, LOCATIONS, load_steps

PREFIX = {'White Orchard': 'WO', 'Visima': 'VZ', 'Velen': 'VEL', 'Novigrad': 'NOV',
          'Skellige': 'SKE', 'Kear Morhen': 'KM', 'Toussaint': 'TOU'}
IN_BLOCK_DOC = 'source is a Block_ID tag, target is a step sealed in it (no free roam until the block is finished)'


def norm(s):
    return re.sub(r'\s+', ' ', str(s)).strip().casefold()


def build(steps):
    uids = set(steps['step_uid'])
    names = collections.defaultdict(set)
    levels = collections.defaultdict(list)       # quest -> [(number, optional, uid)]
    for s in steps.itertuples():
        names[norm(s.quest_name)].add(s.quest_id)
        m = re.match(r'\d+', s.step_id)
        levels[s.quest_id].append((int(m.group()) if m else 0, s.is_optional, s.step_uid))

    def quest_end(qid):
        mand = [x for x in levels[qid] if not x[1]] or levels[qid]
        top = max(x[0] for x in mand)
        return [x[2] for x in mand if x[0] == top]

    def parse_ref(cell):
        """-> (ids, unresolved parts). Grammar: 'A|B' alternatives, 'A/B' or 'QID_S2a/b' branch groups, or a quest name."""
        ids, bad = [], []
        for part in str(cell).split('|'):
            toks = [t.strip() for t in part.split('/')]
            if re.fullmatch(r'[A-Za-z0-9]+_S\d+[a-z]*', toks[0]):
                base = re.match(r'(.*_S\d+)', toks[0]).group(1)
                for t in toks:
                    uid = t if '_S' in t else base + t
                    (ids if uid in uids else bad).append(uid)
            else:
                qs = names.get(norm(part))
                if qs and len(qs) == 1:
                    ids.extend(quest_end(next(iter(qs))))
                else:
                    bad.append(part.strip())
        return ids, bad

    edges = {loc: {} for loc in LOCATIONS}      # loc -> {(src, type, tgt): None} (ordered, deduplicated)
    issues, prereqs = [], collections.defaultdict(set)

    def add(loc, src, typ, tgt):
        if src != tgt:
            edges[loc].setdefault((src, typ, tgt), None)

    for s in steps.itertuples():
        for col, cell, typ in (('After', s.after, None), ('Before', s.before, 'blocks')):
            if cell is None or cell != cell:    # None or NaN
                continue
            ids, bad = parse_ref(cell)
            issues += [(s.location, s.row, s.step_uid, col, str(cell), b) for b in bad]
            if col == 'After':
                prereqs[s.step_uid].update(ids)
            for i in ids:
                add(s.location, i, typ or ('requires' if len(ids) == 1 else 'unlocks'), s.step_uid)
        if pd.notna(s.block_id):
            add(s.location, s.block_id, 'in_block', s.step_uid)

    # Path tags are alternatives of one decision when the quest and the outside prerequisites of their steps match.
    # (Pairing every tag of a quest would make two separate decisions in one quest exclude each other.)
    by_tag = collections.defaultdict(list)
    for s in steps.itertuples():
        for a, b in itertools.combinations(sorted(s.path_tags), 2):     # tags listed together in one cell
            add(s.location, a, 'excludes_path', b)
        if len(s.path_tags) == 1:
            by_tag[s.path_tags[0]].append(s)
    decisions = collections.defaultdict(list)
    for tag, group in by_tag.items():
        own = {s.step_uid for s in group}
        outside = frozenset(p for s in group for p in prereqs[s.step_uid] if p not in own)
        decisions[(group[0].location, group[0].quest_id, outside)].append(tag)
    for (loc, _, _), tags in decisions.items():
        for a, b in itertools.combinations(sorted(tags), 2):
            add(loc, a, 'excludes_path', b)
    return edges, issues


def write(wb, edges, issues):
    for loc in LOCATIONS:
        ws = wb[f'{loc} Relationships']
        for row in ws.iter_rows():
            for c in row:
                c.value = None
        for i, h in enumerate(['relationship_id', 'source_step_id', 'relationship_type', 'target_step_uid'], 1):
            ws.cell(1, i).value = h
            ws.cell(1, i).font = Font(bold=True)
        for n, (src, typ, tgt) in enumerate(edges[loc], 1):
            for i, v in enumerate((f'{PREFIX[loc]}_REL_{n:04d}', src, typ, tgt), 1):
                ws.cell(n + 1, i).value = v
        for col, w in zip('ABCD', (16, 24, 18, 24)):
            ws.column_dimensions[col].width = w

    wr = wb['Relationships']
    if not any(wr.cell(r, 1).value == 'in_block' for r in range(1, wr.max_row + 1)):
        last = max(r for r in range(1, wr.max_row + 1) if wr.cell(r, 1).value)
        wr.cell(last + 1, 1).value, wr.cell(last + 1, 2).value = 'in_block', IN_BLOCK_DOC

    if 'Relationship_issues' in wb.sheetnames:
        del wb['Relationship_issues']
    wi = wb.create_sheet('Relationship_issues')
    for i, h in enumerate(['sheet', 'row', 'step_id', 'column', 'cell_value', 'unresolved_part'], 1):
        wi.cell(1, i).value = h
        wi.cell(1, i).font = Font(bold=True)
    for n, row in enumerate(issues, 2):
        for i, v in enumerate(row, 1):
            wi.cell(n, i).value = v


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--file', default=DATA_FILE)
    args = ap.parse_args()
    wb = openpyxl.load_workbook(args.file)
    edges, issues = build(load_steps(wb))
    write(wb, edges, issues)
    wb.calculation.fullCalcOnLoad = True
    wb.save(args.file)
    for loc in LOCATIONS:
        print(f'{loc:14} {dict(collections.Counter(t for _, t, _ in edges[loc]))}')
    print(f'{len(issues)} unresolved cells (sheet Relationship_issues)')


if __name__ == '__main__':
    main()
