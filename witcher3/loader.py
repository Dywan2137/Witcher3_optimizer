"""Step 1: load the quest workbook (all locations) into tables and validate them."""
from __future__ import annotations

import graphlib
import re
from collections import defaultdict
from dataclasses import dataclass, replace
from pathlib import Path

import openpyxl
import pandas as pd

DATA_FILE = Path(__file__).resolve().parents[1] / 'Data' / 'Witcher_3_quests copy.xlsx'
LOCATIONS = ['White Orchard', 'Visima', 'Velen', 'Novigrad', 'Skellige', 'Kear Morhen', 'Toussaint']
STEP_EDGES = {'requires', 'unlocks', 'blocks', 'excludes', 'during'}   # step -> step
PRECEDENCE = {'requires', 'unlocks'}

SCOPES = {'base': 'podstawa', 'hos': 'Serce z kamienia', 'bw': 'Krew i wino'}   # scope name -> value of the DLC column

HEADER_ALIASES = {
    'describtion': 'description', 'reward_xp': 'reward_xp', 'manditory': 'mandatory',
    'has_level_recommendation': 'has_level_recommendation', 'recommended_level': 'recommended_level',
}


@dataclass
class Dataset:
    steps: pd.DataFrame
    relationships: pd.DataFrame
    reward_rules: pd.DataFrame
    level_table: pd.DataFrame
    rel_types: set


def _header(h) -> str:
    key = re.sub(r'\W+', '_', str(h).strip().lower()).strip('_')
    return HEADER_ALIASES.get(key, key)


def _tag(value, quest_id):
    """None for empty/'-'; evaluates the =CONCAT(D<row>,"suffix") formulas used for Block_ID/Path_ID."""
    if value is None or str(value).strip() in ('', '-'):
        return None
    if isinstance(value, str) and value.startswith('='):
        m = re.fullmatch(r'=_xlfn\.CONCAT\(D\d+,"(.*)"\)', value)
        return quest_id + m.group(1) if m else value
    return str(value).strip()


def split_tags(value) -> tuple:
    """A Path_ID cell may list several mutually exclusive tags: 'WOF4_PATH_A / WOF5_PATH_B / WOF6_PATH_C'."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ()
    return tuple(t.strip() for t in str(value).split('/') if t.strip())


def _ref(value):
    return None if value is None or str(value).strip() in ('', '-') else str(value).strip()


def _bool(value):
    if isinstance(value, bool):
        return value
    return None if value is None else str(value).strip().lower() == 'true'


def _num(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return float('nan')


def load_steps(wb) -> pd.DataFrame:
    rows = []
    for loc in LOCATIONS:
        ws = wb[loc]
        cols = {_header(c.value): c.column for c in ws[1] if c.value is not None}
        for r in range(2, ws.max_row + 1):
            get = lambda k: ws.cell(r, cols[k]).value if k in cols else None
            qid = get('quest_id')
            if not qid:
                continue
            sid = None if get('step_id') is None else str(get('step_id')).strip()
            optional = bool(sid and re.fullmatch(r'\d+[a-z]*o', sid))   # 'o' suffix = optional step
            mandatory = _bool(get('mandatory'))
            path = _tag(get('path_id'), str(qid))
            rows.append(dict(
                location=loc, row=r, step_uid=f'{qid}_S{sid}', quest_id=str(qid).strip(), step_id=sid,
                quest_name=str(get('quest_name') or '').strip(), description=str(get('description') or '').strip(),
                category=get('category'), dlc=get('dlc'), region=get('location'),
                has_level_recommendation=bool(_bool(get('has_level_recommendation'))),
                recommended_level=_num(get('recommended_level')), reward_xp=_num(get('reward_xp')),
                reward_rule_type=get('reward_rule_type'),
                is_optional=optional, mandatory=(not optional) if mandatory is None else mandatory,
                before=_ref(get('before')), after=_ref(get('after')),
                block_id=_tag(get('block_id'), str(qid)), path_id=path, path_tags=split_tags(path),
            ))
    return pd.DataFrame(rows)


def load_relationships(wb) -> pd.DataFrame:
    rows = []
    for loc in LOCATIONS:
        ws = wb[f'{loc} Relationships']
        for r in range(2, ws.max_row + 1):
            v = [ws.cell(r, c).value for c in range(1, 5)]
            if any(x is not None for x in v):
                rows.append(dict(location=loc, row=r, relationship_id=v[0], source=v[1], type=v[2], target=v[3]))
    return pd.DataFrame(rows, columns=['location', 'row', 'relationship_id', 'source', 'type', 'target'])


def load_table(wb, sheet, ncols) -> pd.DataFrame:
    ws = wb[sheet]
    rows = [r for r in ws.iter_rows(min_col=1, max_col=ncols, values_only=True) if any(x is not None for x in r)]
    return pd.DataFrame(rows[1:], columns=[_header(h) for h in rows[0]])


def load(path: Path = DATA_FILE) -> Dataset:
    wb = openpyxl.load_workbook(path)   # formulas, not cached values: the file may have no cached values
    types = load_table(wb, 'Relationships', 2)
    return Dataset(
        steps=load_steps(wb),
        relationships=load_relationships(wb),
        reward_rules=load_table(wb, 'Reward_rules', 5),
        level_table=load_table(wb, 'Level_table', 3),
        rel_types=set(types['type'].dropna().astype(str).str.strip()),
    )


def restrict(ds: Dataset, scope) -> Dataset:
    """Keep only the steps of the given scopes (names from SCOPES, e.g. ('base', 'hos')) and the edges between them.
    A scope may only depend on itself or on earlier scopes: tests/test_data.py checks that base + hos never need Blood and Wine."""
    steps = ds.steps[ds.steps['dlc'].isin({SCOPES[s] for s in scope})].reset_index(drop=True)
    uids, blocks = set(steps['step_uid']), set(steps['block_id'].dropna())
    tags = {t for ts in steps['path_tags'] for t in ts}
    rel = ds.relationships
    ends = {'in_block': (blocks, uids), 'excludes_path': (tags, tags)}
    keep = [e.source in ends.get(e.type, (uids, uids))[0] and e.target in ends.get(e.type, (uids, uids))[1]
            for e in rel.itertuples()]
    return replace(ds, steps=steps, relationships=rel[keep].reset_index(drop=True))


def edges_leaving_scope(ds: Dataset, scope) -> pd.DataFrame:
    """Prerequisite edges (requires/unlocks) from a step outside `scope` into a step inside it: the in-scope step
    would silently lose that prerequisite under restrict(). Should be empty."""
    inside = set(restrict(ds, scope).steps['step_uid'])
    r = ds.relationships
    return r[r['type'].isin(PRECEDENCE) & r['target'].isin(inside) & ~r['source'].isin(inside)]


# ----------------------------------------------------------------------------------------------
# validation
# ----------------------------------------------------------------------------------------------
def validate(ds: Dataset) -> pd.DataFrame:
    issues = []

    def add(sev, check, location, ref, msg):
        issues.append(dict(severity=sev, check=check, location=location, ref=ref, message=msg))

    st, rel = ds.steps, ds.relationships
    uids = set(st['step_uid'])
    loc_of = dict(zip(st['step_uid'], st['location']))

    # --- steps ---
    for uid, g in st.groupby('step_uid'):
        if len(g) > 1:
            add('error', 'duplicate_step_uid', g.iloc[0]['location'], uid, f'rows {list(g["row"])}')
    for _, s in st.iterrows():
        if not s['step_id'] or s['step_id'] == 'None':
            add('error', 'missing_step_id', s['location'], f'row {s["row"]}', s['quest_id'])
        if pd.isna(s['reward_xp']) or s['reward_xp'] < 0:
            add('error', 'bad_reward_xp', s['location'], s['step_uid'], f'reward_xp={s["reward_xp"]}')
        if s['has_level_recommendation'] and pd.isna(s['recommended_level']):
            add('error', 'missing_recommended_level', s['location'], s['step_uid'], 'has_level_recommendation but no level')
        if s['reward_rule_type'] not in set(ds.reward_rules['reward_rule_type']):
            add('error', 'unknown_reward_rule', s['location'], s['step_uid'],
                'empty' if pd.isna(s['reward_rule_type']) else repr(s['reward_rule_type']))

    # --- reward rules and level table ---
    for typ, g in ds.reward_rules.groupby('reward_rule_type'):
        g = g.sort_values('level_gap_min')
        for (_, a), (_, b) in zip(g.iterrows(), g.iloc[1:].iterrows()):
            if b['level_gap_min'] != a['level_gap_max'] + 1:
                add('error', 'reward_rule_gap_or_overlap', 'Reward_rules', typ,
                    f'{a["level_gap_min"]}..{a["level_gap_max"]} then {b["level_gap_min"]}..{b["level_gap_max"]}')
    lt = ds.level_table
    if list(lt['level']) != list(range(1, len(lt) + 1)):
        add('error', 'level_table_levels', 'Level_table', '', 'levels are not 1..N')
    cum = list(lt['cumulative_xp_to_reach_level'])
    for i in range(len(lt) - 1):
        if cum[i + 1] - cum[i] != lt['xp_to_next_level'].iloc[i]:
            add('error', 'level_table_inconsistent', 'Level_table', f'level {i + 1}',
                f'cumulative step {cum[i + 1] - cum[i]} != xp_to_next_level {lt["xp_to_next_level"].iloc[i]}')

    # --- relationships ---
    block_tags = set(st['block_id'].dropna())
    path_tags = {t for tags in st['path_tags'] for t in tags}
    for _, e in rel.iterrows():
        where = (e['location'], f'{e["relationship_id"]} (row {e["row"]})')
        if pd.isna(e['source']) or pd.isna(e['type']) or pd.isna(e['target']):
            add('error', 'incomplete_relationship', *where, 'empty source/type/target')
            continue
        if e['type'] not in ds.rel_types:
            add('error', 'unknown_relationship_type', *where, e['type'])
        if e['source'] == e['target']:
            add('error', 'self_edge', *where, e['source'])
        ends = {'in_block': (block_tags, uids), 'excludes_path': (path_tags, path_tags)}.get(
            e['type'], (uids, uids) if e['type'] in STEP_EDGES else (None, None))
        for end, pool in zip(('source', 'target'), ends):
            if pool is not None and e[end] not in pool:
                add('error', 'dangling_endpoint', *where, f'{end} {e[end]!r} not found')
    dup = rel[rel.duplicated(['source', 'type', 'target'], keep='first')]
    for _, e in dup.iterrows():
        add('warning', 'duplicate_relationship', e['location'], e['relationship_id'], f'{e["source"]} {e["type"]} {e["target"]}')

    steps_rel = rel[rel['type'].isin(STEP_EDGES) & rel['source'].isin(uids) & rel['target'].isin(uids)]
    req, unl, blk = defaultdict(set), defaultdict(set), []
    for _, e in steps_rel.iterrows():
        if e['type'] == 'requires':
            req[e['target']].add(e['source'])
        elif e['type'] == 'unlocks':
            unl[e['target']].add(e['source'])
        elif e['type'] == 'blocks':
            blk.append((e['source'], e['target']))

    # cycles in the precedence graph
    ts = graphlib.TopologicalSorter()
    for t in uids:
        ts.add(t, *(req[t] | unl[t]))
    try:
        ts.prepare()
    except graphlib.CycleError as err:
        cyc = err.args[1]
        add('error', 'precedence_cycle', loc_of.get(cyc[0], ''), ' -> '.join(cyc), 'requires/unlocks form a cycle')

    # steps that can never become available (ignoring blocks/exclusions): all requires done, any unlocks done
    avail, changed = set(), True
    while changed:
        changed = False
        for t in uids - avail:
            if all(s in avail for s in req[t]) and (not unl[t] or any(s in avail for s in unl[t])):
                avail.add(t)
                changed = True
    for t in sorted(uids - avail):
        add('error', 'unreachable_step', loc_of[t], t, 'prerequisites can never be satisfied')

    # a blocker that is also a (transitive) prerequisite of its own target
    parents ={t: req[t] | unl[t] for t in uids}
    for s, t in blk:
        seen, stack = set(), [t]
        while stack:
            for p in parents[stack.pop()]:
                if p not in seen:
                    seen.add(p)
                    stack.append(p)
        if s in seen:
            add('error', 'blocker_is_prerequisite', loc_of[t], f'{s} blocks {t}', 'source is required to reach the target')

    # source cells that did not become edges (unresolved names / unknown ids)
    has_in = {k: set(g['target']) for k, g in steps_rel.groupby('type')}
    for _, s in st.iterrows():
        if pd.notna(s['after']) and s['step_uid'] not in has_in.get('requires', set()) | has_in.get('unlocks', set()):
            add('warning', 'after_without_edge', s['location'], s['step_uid'], f'After={s["after"]!r}')
        if pd.notna(s['before']) and s['step_uid'] not in has_in.get('blocks', set()):
            add('warning', 'before_without_edge', s['location'], s['step_uid'], f'Before={s["before"]!r}')

    # Block_ID / Path_ID versus the generated relationship sheets, plus tag sanity
    in_block = rel[rel['type'] == 'in_block']
    have = set(zip(in_block['source'], in_block['target']))
    want = set(zip(st['block_id'].dropna(), st.loc[st['block_id'].notna(), 'step_uid']))
    for src, tgt in sorted(want - have):
        add('warning', 'relationships_out_of_sync', loc_of[tgt], tgt, f'Block_ID {src} has no in_block edge')
    for src, tgt in sorted(have - want):
        add('warning', 'relationships_out_of_sync', loc_of.get(tgt, ''), tgt, f'in_block edge {src} but no Block_ID in the sheet')
    # tags declared as alternatives of each other by a shared anchor step (quest-level choices such as KIWF11 / KIWF12)
    alternatives = {t for tags in st['path_tags'] if len(tags) > 1 for t in tags}
    for s in st.itertuples():                  # branch letters and path tags should go together
        m = re.fullmatch(r'\d+([a-z]*)', str(s.step_id))
        letter = m.group(1).removesuffix('o') if m else ''
        if letter and not s.path_tags:
            add('warning', 'branch_letter_without_path_tag', s.location, s.step_uid, 'alternative step has no Path_ID')
        if len(s.path_tags) == 1 and not letter and s.path_tags[0] not in alternatives:
            add('warning', 'path_tag_on_unlettered_step', s.location, s.step_uid, f'{s.path_tags[0]} on a step with no branch letter')
    by_tag = defaultdict(list)
    for s in st.itertuples():
        if len(s.path_tags) == 1:
            by_tag[s.path_tags[0]].append(s.step_uid)
    for tag, members in by_tag.items():        # one tag should be one chain of steps, i.e. one decision
        parent = {u: u for u in members}

        def find(x):
            while parent[x] != x:
                x = parent[x]
            return x
        for u in members:
            for p in req[u] | unl[u]:
                if p in parent:
                    parent[find(u)] = find(p)
        if len({find(u) for u in members}) > 1:
            add('warning', 'path_tag_reused', loc_of[members[0]], tag, 'tag covers separate chains: two decisions share one tag')

    paired = set(rel.loc[rel['type'] == 'excludes_path', 'source']) | set(rel.loc[rel['type'] == 'excludes_path', 'target'])
    for tag in sorted(path_tags - paired):
        first = st[st['path_tags'].apply(lambda t: tag in t)].iloc[0]
        add('warning', 'path_without_partner', first['location'], first['quest_id'], f'{tag} has no excludes_path edge')

    out = pd.DataFrame(issues, columns=['severity', 'check', 'location', 'ref', 'message'])
    return out.sort_values(['severity', 'check', 'location'], ignore_index=True)


def summarize(ds: Dataset) -> pd.DataFrame:
    st, rel = ds.steps, ds.relationships
    g = st.groupby('location', sort=False)
    out = pd.DataFrame({
        'steps': g.size(), 'quests': g['quest_id'].nunique(), 'total_xp': g['reward_xp'].sum(),
        'optional_steps': g['is_optional'].sum(), 'block_tags': g['block_id'].nunique(),
        'path_tags': g['path_tags'].agg(lambda s: len({t for tags in s for t in tags})),
    })
    edges = rel.pivot_table(index='location', columns='type', values='row', aggfunc='count', fill_value=0)
    edges.columns = [f'edges_{c}' for c in edges.columns]
    return out.join(edges).fillna(0).astype(int)
