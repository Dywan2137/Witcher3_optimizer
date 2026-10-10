"""Write the best playthrough to an Excel file the player can follow.

Sheets
  Plan     one row per step in the order to play it: all the original columns except Before / After / Block_ID /
           Path_ID (those only matter to the search), then the player's level before the step, the multiplier it paid,
           the XP earned, the running XP and the level after. Rows where the level goes up are highlighted.
  Skipped  the steps the plan does not play and why (another path was chosen, a story step cuts them off, ...).
  Summary  scope, final level, XP, and the step at which each level is reached.

Usage: python -m witcher3.export [--scope base,hos,bw] [--restarts 25] [--noise 2] [--out output/optimal_order.xlsx]
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

if __package__ in (None, ''):  # run as a plain script (python witcher3/export.py, an IDE Run button)
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from witcher3.baselines import Run, lost_steps
from witcher3.greedy import CONFIGS, best_of
from witcher3.loader import DATA_FILE, load, restrict
from witcher3.simulator import Model, is_dead
from witcher3.units import Plan

DEFAULT_OUT = Path(__file__).resolve().parents[1] / 'output' / 'optimal_order.xlsx'
FONT = 'Arial'

# original columns kept in the player's file, as (header, key in the loader's step table)
ORIGINAL = [('Category', 'category'), ('DLC', 'dlc'), ('ID', 'step_uid'), ('Quest_ID', 'quest_id'), ('Step_ID', 'step_id'),
            ('Quest_Name', 'quest_name'), ('Description', 'description'), ('Location', 'region'),
            ('Has_level_recommendation', 'has_level_recommendation'), ('Recommended_level', 'recommended_level'),
            ('Reward XP', 'reward_xp'), ('Mandatory', 'mandatory'), ('Reward_rule_type', 'reward_rule_type')]
PLAN_HEAD = ['Order', 'Area'] + [h for h, _ in ORIGINAL] + ['Level before', 'Multiplier', 'XP earned', 'Total XP', 'Level after']
WIDTHS = {'Order': 7, 'Area': 14, 'Category': 20, 'DLC': 16, 'ID': 13, 'Quest_ID': 11, 'Step_ID': 8, 'Quest_Name': 34,
          'Description': 60, 'Location': 16, 'Has_level_recommendation': 12, 'Recommended_level': 12, 'Reward XP': 10,
          'Mandatory': 10, 'Reward_rule_type': 15, 'Level before': 9, 'Multiplier': 10, 'XP earned': 10, 'Total XP': 10,
          'Level after': 9}


def cell_value(key, value):
    if key == 'step_id' and str(value).isdigit():
        return int(value)                       # the original sheet stores plain numbers as numbers
    if isinstance(value, float) and value != value:
        return None                             # NaN
    if hasattr(value, 'item'):
        return value.item()
    return value


def skipped_rows(plan: Plan, ds, run: Run):
    m, s = plan.model, run.state
    info = ds.steps.set_index('step_uid')
    played = {m.idx[u] for u in run.order}
    lost = lost_steps(m, s, played)
    out = []
    for i in range(len(m.uids)):
        if i in played:
            continue
        if m.tags[i] and all(t in s.excluded for t in m.tags[i]):
            why = 'Another path of this quest was chosen'
        elif i in s.dead:
            why = 'Cut off by an earlier story step or exclusive choice'
        elif i in lost:
            why = 'Needs a step that is skipped'
        else:
            why = 'Never became available (check the data)'
        out.append((i, why))
    return [(info.loc[m.uids[i]], why) for i, why in out]


def style_header(ws, row, ncols, fill='1F3864'):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', start_color=fill)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def write_plan(wb, ds, run: Run):
    ws = wb.create_sheet('Plan')
    ws.append(PLAN_HEAD)
    style_header(ws, 1, len(PLAN_HEAD))
    info = ds.steps.set_index('step_uid')
    thin, up = Side(style='thin', color='999999'), PatternFill('solid', start_color='FFF2CC')
    last_quest = None
    for n, t in enumerate(run.trace, 1):
        s = info.loc[t.uid]
        row = [n, s['location']] + [cell_value(k, t.uid if k == 'step_uid' else s[k]) for _, k in ORIGINAL] + \
              [t.level_before, t.multiplier, round(t.gain, 1), round(t.xp), t.level]
        ws.append(row)
        r = ws.max_row
        for c in range(1, len(PLAN_HEAD) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name=FONT, bold=(PLAN_HEAD[c - 1] in ('Level after',) and t.level > t.level_before))
            if t.level > t.level_before:
                cell.fill = up
            if s['quest_id'] != last_quest:
                cell.border = Border(top=thin)
        ws.cell(r, PLAN_HEAD.index('Multiplier') + 1).number_format = '0%'
        ws.cell(r, PLAN_HEAD.index('XP earned') + 1).number_format = '#,##0.0'
        ws.cell(r, PLAN_HEAD.index('Total XP') + 1).number_format = '#,##0'
        last_quest = s['quest_id']
    for i, h in enumerate(PLAN_HEAD, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[h]
    ws.freeze_panes = ws.cell(2, PLAN_HEAD.index('Description') + 1)
    ws.auto_filter.ref = f'A1:{get_column_letter(len(PLAN_HEAD))}{ws.max_row}'
    return ws


def write_skipped(wb, plan: Plan, ds, run: Run):
    ws = wb.create_sheet('Skipped')
    head = ['Area', 'Category', 'DLC', 'ID', 'Quest_ID', 'Step_ID', 'Quest_Name', 'Description', 'Reward XP', 'Why skipped']
    ws.append(head)
    style_header(ws, 1, len(head))
    for s, why in skipped_rows(plan, ds, run):
        ws.append([s['location'], s['category'], s['dlc'], s.name, s['quest_id'], cell_value('step_id', s['step_id']),
                   s['quest_name'], s['description'], cell_value('x', s['reward_xp']), why])
        for c in range(1, len(head) + 1):
            ws.cell(ws.max_row, c).font = Font(name=FONT)
    for i, w in enumerate([14, 20, 16, 13, 11, 8, 34, 60, 10, 52], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(head))}{ws.max_row}'


def write_summary(wb, plan: Plan, ds, run: Run, scope, level_table):
    ws = wb.create_sheet('Summary', 0)
    cum = list(level_table['cumulative_xp_to_reach_level'])
    skipped = len(plan.model.uids) - len(run.trace)
    facts = [('Scope', ', '.join(scope)), ('Steps to play', len(run.trace)), ('Steps skipped', skipped),
             ('Final level', run.level), ('Total quest XP', round(run.xp)),
             ('XP needed for the next level', cum[run.level] if run.level < len(cum) else 'max level')]
    ws.append(['Optimal quest order']); ws['A1'].font = Font(name=FONT, bold=True, size=14)
    ws.append([])
    for k, v in facts:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(name=FONT, bold=True)
        ws.cell(ws.max_row, 2).font = Font(name=FONT)
        ws.cell(ws.max_row, 2).alignment = Alignment(horizontal='left')
    ws.append([])
    notes = ['Starts at level 1 and counts quest XP only (no monster nests, abandoned sites, combat or Gwent XP).',
             'XP per step = Reward XP x multiplier from the Reward_rules sheet (Death March), set by the level before the step.',
             'Follow the Plan sheet from top to bottom. Highlighted rows are the steps where you level up.',
             'The Skipped sheet lists what the plan leaves out and why.']
    for n in notes:
        ws.append([n]); ws.cell(ws.max_row, 1).font = Font(name=FONT, italic=True)
    ws.append([])
    ws.append(['Level', 'Reached at step', 'Step', 'Quest'])
    style_header(ws, ws.max_row, 4)
    info = ds.steps.set_index('step_uid')
    seen = set()
    for n, t in enumerate(run.trace, 1):
        if t.level not in seen and t.level > 1:
            seen.add(t.level)
            ws.append([t.level, n, t.uid, info.loc[t.uid]['quest_name']])
            for c in range(1, 5):
                ws.cell(ws.max_row, c).font = Font(name=FONT)
    for col, w in zip('ABCD', (30, 16, 16, 40)):
        ws.column_dimensions[col].width = w


def export(ds, plan: Plan, run: Run, path, scope, level_table) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_plan(wb, ds, run)
    write_skipped(wb, plan, ds, run)
    write_summary(wb, plan, ds, run, scope, level_table)
    wb.save(path)
    return path


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('--file', default=DATA_FILE)
    ap.add_argument('--scope', default='base,hos,bw', help='comma separated: base, hos, bw')
    ap.add_argument('--restarts', type=int, default=25, help='randomized restarts per rule setting (0 = plain greedy only)')
    ap.add_argument('--noise', type=float, default=2.0)
    ap.add_argument('--out', default=str(DEFAULT_OUT))
    args = ap.parse_args()

    scope = args.scope.split(',')
    full = load(args.file)
    ds = restrict(full, scope)
    plan = Plan(Model(ds))
    run = best_of(plan, args.restarts, args.noise, configs=CONFIGS if args.restarts else None)
    out = export(ds, plan, run, args.out, scope, full.level_table)
    print(f'scope {args.scope}: level {run.level}, {run.xp:,.0f} XP, {len(run.trace)} steps -> {out}')


if __name__ == '__main__':
    main()
